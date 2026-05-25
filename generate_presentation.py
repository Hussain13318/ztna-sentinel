from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt


OUT_PATH = Path(__file__).resolve().parent / "ZTNA_Framework_Presentation.pptx"

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)

# Colors
BG = RGBColor(15, 23, 42)
BG_2 = RGBColor(30, 41, 59)
ACCENT = RGBColor(56, 189, 248)
ACCENT_2 = RGBColor(14, 165, 233)
TEXT = RGBColor(248, 250, 252)
MUTED = RGBColor(203, 213, 225)
GREEN = RGBColor(34, 197, 94)
YELLOW = RGBColor(250, 204, 21)
RED = RGBColor(239, 68, 68)


def set_bg(slide, color=BG):
    shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, prs.slide_height)
    shape.fill.solid()
    shape.fill.fore_color.rgb = color
    shape.line.fill.background()
    slide.shapes._spTree.remove(shape._element)
    slide.shapes._spTree.insert(2, shape._element)


def add_top_bar(slide, title, subtitle=None):
    bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, Inches(0.8))
    bar.fill.solid()
    bar.fill.fore_color.rgb = BG_2
    bar.line.fill.background()

    title_box = slide.shapes.add_textbox(Inches(0.55), Inches(0.18), Inches(10.2), Inches(0.35))
    tf = title_box.text_frame
    p = tf.paragraphs[0]
    r = p.add_run()
    r.text = title
    r.font.size = Pt(24)
    r.font.bold = True
    r.font.color.rgb = TEXT

    if subtitle:
        sub = slide.shapes.add_textbox(Inches(10.6), Inches(0.2), Inches(2.0), Inches(0.3))
        tf2 = sub.text_frame
        p2 = tf2.paragraphs[0]
        p2.alignment = PP_ALIGN.RIGHT
        r2 = p2.add_run()
        r2.text = subtitle
        r2.font.size = Pt(11)
        r2.font.color.rgb = MUTED


def add_footer(slide, text="ZTNA Framework | Air University"):
    footer = slide.shapes.add_textbox(Inches(0.55), Inches(7.05), Inches(6.5), Inches(0.25))
    tf = footer.text_frame
    p = tf.paragraphs[0]
    r = p.add_run()
    r.text = text
    r.font.size = Pt(10)
    r.font.color.rgb = MUTED


def add_bullets(slide, left, top, width, height, items, font_size=22, color=TEXT, bullet_color=ACCENT):
    box = slide.shapes.add_textbox(left, top, width, height)
    tf = box.text_frame
    tf.word_wrap = True
    tf.margin_left = 0
    tf.margin_right = 0
    tf.margin_top = 0
    tf.margin_bottom = 0
    first = True
    for text, level in items:
        p = tf.paragraphs[0] if first else tf.add_paragraph()
        first = False
        p.level = level
        p.space_after = Pt(8)
        p.line_spacing = 1.1
        r = p.add_run()
        r.text = f"- {text}"
        r.font.size = Pt(font_size if level == 0 else font_size - 2)
        r.font.color.rgb = color
    return box


def add_card(slide, left, top, width, height, title, body, accent=ACCENT):
    card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
    card.fill.solid()
    card.fill.fore_color.rgb = BG_2
    card.line.color.rgb = accent
    card.line.width = Pt(1.5)

    tbox = slide.shapes.add_textbox(left + Inches(0.18), top + Inches(0.12), width - Inches(0.36), Inches(0.35))
    tf = tbox.text_frame
    p = tf.paragraphs[0]
    r = p.add_run()
    r.text = title
    r.font.size = Pt(16)
    r.font.bold = True
    r.font.color.rgb = TEXT

    bbox = slide.shapes.add_textbox(left + Inches(0.18), top + Inches(0.5), width - Inches(0.36), height - Inches(0.62))
    tfb = bbox.text_frame
    tfb.word_wrap = True
    p2 = tfb.paragraphs[0]
    p2.alignment = PP_ALIGN.LEFT
    r2 = p2.add_run()
    r2.text = body
    r2.font.size = Pt(12)
    r2.font.color.rgb = MUTED


# Slide 1: Title
slide = prs.slides.add_slide(prs.slide_layouts[6])
set_bg(slide)
shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, Inches(1.0))
shape.fill.solid()
shape.fill.fore_color.rgb = BG_2
shape.line.fill.background()

accent = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(6.85), prs.slide_width, Inches(0.65))
accent.fill.solid()
accent.fill.fore_color.rgb = ACCENT_2
accent.line.fill.background()

box = slide.shapes.add_textbox(Inches(0.75), Inches(1.2), Inches(11.6), Inches(2.0))
tf = box.text_frame
p = tf.paragraphs[0]
r = p.add_run()
r.text = "AI-Based Adaptive Hybrid Zero Trust Network Access Framework"
r.font.size = Pt(28)
r.font.bold = True
r.font.color.rgb = TEXT

p2 = tf.add_paragraph()
p2.space_before = Pt(12)
r2 = p2.add_run()
r2.text = "Presentation: Introduction, Literature Review, Proposed Approach, Results, Zero Trust Aspect, and Future Work"
r2.font.size = Pt(16)
r2.font.color.rgb = MUTED

box2 = slide.shapes.add_textbox(Inches(0.75), Inches(3.5), Inches(8.5), Inches(1.0))
tf2 = box2.text_frame
p3 = tf2.paragraphs[0]
r3 = p3.add_run()
r3.text = "ZTNA security model with MFA, device posture checks, adaptive risk scoring, RBAC, and session logging."
r3.font.size = Pt(18)
r3.font.color.rgb = TEXT

meta = slide.shapes.add_textbox(Inches(0.75), Inches(6.9), Inches(12.0), Inches(0.35))
mtf = meta.text_frame
mp = mtf.paragraphs[0]
mr = mp.add_run()
mr.text = "Air University | 4th Semester Project"
mr.font.size = Pt(11)
mr.font.color.rgb = TEXT

# Slide 2: Introduction
slide = prs.slides.add_slide(prs.slide_layouts[6])
set_bg(slide)
add_top_bar(slide, "Introduction", "01")
add_bullets(
    slide,
    Inches(0.7), Inches(1.2), Inches(6.1), Inches(5.5),
    [
        ("Traditional perimeter security assumes trust inside the network, which is no longer safe.", 0),
        ("This project implements a Zero Trust Network Access style workflow for user access decisions.", 0),
        ("Every session is evaluated using identity, MFA, device posture, risk score, and role.", 0),
        ("The system is designed as a security demo and academic prototype, not a production gateway.", 0),
    ],
    font_size=20,
)
add_card(slide, Inches(7.1), Inches(1.3), Inches(5.6), Inches(1.4), "Problem", "A login alone is not enough to decide access when devices, users, and network conditions change.")
add_card(slide, Inches(7.1), Inches(2.95), Inches(5.6), Inches(1.4), "Goal", "Make access decisions adaptive, explainable, and logged for audit and review.", accent=GREEN)
add_card(slide, Inches(7.1), Inches(4.6), Inches(5.6), Inches(1.4), "Scope", "Desktop GUI plus a lightweight dashboard showing sessions, alerts, and traffic monitoring.", accent=YELLOW)
add_footer(slide)

# Slide 3: Literature
slide = prs.slides.add_slide(prs.slide_layouts[6])
set_bg(slide)
add_top_bar(slide, "Literature Review", "02")
add_card(slide, Inches(0.7), Inches(1.25), Inches(3.95), Inches(2.0), "Zero Trust Principles", "NIST-style Zero Trust guidance says to never assume trust, verify explicitly, and apply least privilege continuously.", accent=ACCENT)
add_card(slide, Inches(4.67), Inches(1.25), Inches(3.95), Inches(2.0), "Context-Aware Access", "Research and industry practice increasingly use device posture, location, and behavior to make runtime access decisions.", accent=GREEN)
add_card(slide, Inches(8.64), Inches(1.25), Inches(3.95), Inches(2.0), "Risk-Based Auth", "Adaptive authentication combines MFA with risk scoring to tighten access when signals become suspicious.", accent=YELLOW)
add_bullets(
    slide,
    Inches(0.8), Inches(3.65), Inches(12.0), Inches(2.55),
    [
        ("The literature consistently moves from static perimeter defense to continuous verification.", 0),
        ("Zero Trust architectures favor identity, device health, and session context over location alone.", 0),
        ("This project follows that direction by combining MFA, posture checks, RBAC, and session monitoring.", 0),
    ],
    font_size=18,
)
add_footer(slide)

# Slide 4: Proposed approach
slide = prs.slides.add_slide(prs.slide_layouts[6])
set_bg(slide)
add_top_bar(slide, "Proposed Approach", "03")
add_card(slide, Inches(0.7), Inches(1.25), Inches(4.0), Inches(2.1), "1. Authenticate", "User enters username, password, and role. OTP-based MFA is triggered before access continues.", accent=ACCENT)
add_card(slide, Inches(4.67), Inches(1.25), Inches(4.0), Inches(2.1), "2. Evaluate Context", "Device posture, IP context, and role are checked and fed into a dynamic risk scorer.", accent=GREEN)
add_card(slide, Inches(8.64), Inches(1.25), Inches(4.0), Inches(2.1), "3. Decide Access", "Access is granted or denied based on the resulting risk level and policy logic.", accent=RED)
add_card(slide, Inches(0.9), Inches(3.8), Inches(5.8), Inches(2.0), "Core Modules", "MFA, device posture checker, risk engine, RBAC, anomaly detection, traffic monitor, threat response, and session manager.", accent=ACCENT_2)
add_card(slide, Inches(6.95), Inches(3.8), Inches(5.65), Inches(2.0), "Outputs", "Risk score, access decision, session history, alerts, and live traffic statistics in the GUI.", accent=YELLOW)
add_footer(slide)

# Slide 5: Zero Trust aspect
slide = prs.slides.add_slide(prs.slide_layouts[6])
set_bg(slide)
add_top_bar(slide, "Zero Trust Aspect", "04")
add_bullets(
    slide,
    Inches(0.75), Inches(1.25), Inches(6.1), Inches(5.2),
    [
        ("Never trust by default: every login request is re-validated.", 0),
        ("Least privilege: role-based access control limits what each user can do.", 0),
        ("Continuous verification: risk is recalculated using posture and behavioral signals.", 0),
        ("Assume breach: monitoring, alerts, and session blocking are built into the workflow.", 0),
        ("Visibility and auditability: all sessions are stored in JSON and exported to CSV.", 0),
    ],
    font_size=20,
)
add_card(slide, Inches(7.15), Inches(1.35), Inches(5.45), Inches(1.35), "Identity", "Username, password, role, and OTP.", accent=ACCENT)
add_card(slide, Inches(7.15), Inches(2.95), Inches(5.45), Inches(1.35), "Device", "Posture and health checks.", accent=GREEN)
add_card(slide, Inches(7.15), Inches(4.55), Inches(5.45), Inches(1.35), "Session", "Risk score, monitoring, and blocking.", accent=YELLOW)
add_footer(slide)

# Slide 6: Results
slide = prs.slides.add_slide(prs.slide_layouts[6])
set_bg(slide)
add_top_bar(slide, "Results", "05")
add_card(slide, Inches(0.7), Inches(1.25), Inches(4.0), Inches(2.1), "Working GUI", "A Tkinter-based interface was implemented with login, active sessions, alerts, traffic, and history tabs.", accent=ACCENT)
add_card(slide, Inches(4.67), Inches(1.25), Inches(4.0), Inches(2.1), "Adaptive Decisions", "The workflow produces a risk score and marks access as granted or denied depending on the evaluation.", accent=GREEN)
add_card(slide, Inches(8.64), Inches(1.25), Inches(4.0), Inches(2.1), "Session Visibility", "Sessions are logged, replayable, exportable, and can be blocked from the interface.", accent=YELLOW)
add_bullets(
    slide,
    Inches(0.8), Inches(3.8), Inches(12.0), Inches(2.2),
    [
        ("MFA, posture checking, and risk scoring are integrated into a single access flow.", 0),
        ("Traffic monitoring and anomaly detection support security visibility during a session.", 0),
        ("The project demonstrates the practical shape of Zero Trust in a small academic system.", 0),
    ],
    font_size=18,
)
add_footer(slide)

# Slide 7: Future work
slide = prs.slides.add_slide(prs.slide_layouts[6])
set_bg(slide)
add_top_bar(slide, "Future Work", "06")
add_bullets(
    slide,
    Inches(0.75), Inches(1.25), Inches(6.2), Inches(5.3),
    [
        ("Add stronger machine learning models for anomaly detection and behavior profiling.", 0),
        ("Connect to a real identity provider and centralized policy engine.", 0),
        ("Use a database instead of flat files for scalable session and audit storage.", 0),
        ("Add dashboards for security analytics and policy enforcement metrics.", 0),
        ("Improve packet capture and threat response automation for live environments.", 0),
    ],
    font_size=20,
)
add_card(slide, Inches(7.15), Inches(1.4), Inches(5.45), Inches(1.6), "Presentation Tip", "Emphasize that the project is a working prototype that maps well to Zero Trust concepts and can be extended into a larger enterprise design.", accent=ACCENT_2)
add_card(slide, Inches(7.15), Inches(3.3), Inches(5.45), Inches(1.6), "Closing Line", "The project shows how Zero Trust can be implemented as an adaptive access workflow rather than a single authentication event.", accent=GREEN)
add_footer(slide)

prs.save(str(OUT_PATH))
print(f"Saved presentation to {OUT_PATH}")
