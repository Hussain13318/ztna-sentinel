from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.util import Inches, Pt

OUT_PATH = Path(__file__).resolve().parent / "ZTNA_5_Slides_Presentation.pptx"

# Update these names as needed.
TEAM_NAMES = [
    "Your Name 1",
    "Your Name 2",
    "Your Name 3",
]

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)

BG = RGBColor(16, 24, 40)
CARD = RGBColor(30, 41, 59)
TEXT = RGBColor(248, 250, 252)
MUTED = RGBColor(203, 213, 225)
ACCENT = RGBColor(14, 165, 233)


def set_background(slide):
    bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, prs.slide_height)
    bg.fill.solid()
    bg.fill.fore_color.rgb = BG
    bg.line.fill.background()
    slide.shapes._spTree.remove(bg._element)
    slide.shapes._spTree.insert(2, bg._element)


def add_title(slide, title, subtitle=""):
    bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, Inches(0.85))
    bar.fill.solid()
    bar.fill.fore_color.rgb = CARD
    bar.line.fill.background()

    tbox = slide.shapes.add_textbox(Inches(0.6), Inches(0.18), Inches(10.5), Inches(0.45))
    tf = tbox.text_frame
    p = tf.paragraphs[0]
    r = p.add_run()
    r.text = title
    r.font.size = Pt(26)
    r.font.bold = True
    r.font.color.rgb = TEXT

    if subtitle:
        sbox = slide.shapes.add_textbox(Inches(10.9), Inches(0.23), Inches(2.1), Inches(0.3))
        stf = sbox.text_frame
        sp = stf.paragraphs[0]
        sr = sp.add_run()
        sr.text = subtitle
        sr.font.size = Pt(12)
        sr.font.color.rgb = MUTED


def add_bullets(slide, x, y, w, h, lines, size=20):
    box = slide.shapes.add_textbox(x, y, w, h)
    tf = box.text_frame
    tf.word_wrap = True
    first = True
    for line in lines:
        p = tf.paragraphs[0] if first else tf.add_paragraph()
        first = False
        run = p.add_run()
        run.text = f"- {line}"
        run.font.size = Pt(size)
        run.font.color.rgb = TEXT


def add_card(slide, x, y, w, h, heading, body):
    card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, y, w, h)
    card.fill.solid()
    card.fill.fore_color.rgb = CARD
    card.line.color.rgb = ACCENT
    card.line.width = Pt(1.5)

    hbox = slide.shapes.add_textbox(x + Inches(0.2), y + Inches(0.15), w - Inches(0.4), Inches(0.35))
    htf = hbox.text_frame
    hp = htf.paragraphs[0]
    hr = hp.add_run()
    hr.text = heading
    hr.font.size = Pt(16)
    hr.font.bold = True
    hr.font.color.rgb = TEXT

    bbox = slide.shapes.add_textbox(x + Inches(0.2), y + Inches(0.55), w - Inches(0.4), h - Inches(0.75))
    btf = bbox.text_frame
    bp = btf.paragraphs[0]
    br = bp.add_run()
    br.text = body
    br.font.size = Pt(13)
    br.font.color.rgb = MUTED


# Slide 1: Intro + names
slide = prs.slides.add_slide(prs.slide_layouts[6])
set_background(slide)
add_title(slide, "AI-Based Adaptive Hybrid ZTNA Framework", "Slide 1")

intro = slide.shapes.add_textbox(Inches(0.8), Inches(1.4), Inches(12.0), Inches(1.6))
itf = intro.text_frame
ip = itf.paragraphs[0]
ir = ip.add_run()
ir.text = "Zero Trust Network Access Project Presentation"
ir.font.size = Pt(30)
ir.font.bold = True
ir.font.color.rgb = TEXT

name_text = "Presented by: " + ", ".join(TEAM_NAMES)
nb = slide.shapes.add_textbox(Inches(0.8), Inches(3.4), Inches(11.8), Inches(0.8))
ntf = nb.text_frame
np = ntf.paragraphs[0]
nr = np.add_run()
nr.text = name_text
nr.font.size = Pt(20)
nr.font.color.rgb = MUTED

course = slide.shapes.add_textbox(Inches(0.8), Inches(4.25), Inches(11.5), Inches(0.6))
ctf = course.text_frame
cp = ctf.paragraphs[0]
cr = cp.add_run()
cr.text = "Air University | 4th Semester"
cr.font.size = Pt(16)
cr.font.color.rgb = MUTED

# Slide 2: Introduction
slide = prs.slides.add_slide(prs.slide_layouts[6])
set_background(slide)
add_title(slide, "Introduction", "Slide 2")
add_bullets(
    slide,
    Inches(0.8),
    Inches(1.3),
    Inches(12.0),
    Inches(4.8),
    [
        "Traditional security trusts users after login, which is risky in modern networks.",
        "Zero Trust verifies users and devices continuously before and during access.",
        "Our project combines MFA, device posture checks, risk scoring, and role-based access control.",
        "Goal: smarter, adaptive access decisions with session visibility and threat response.",
    ],
    size=21,
)

# Slide 3: Literature
slide = prs.slides.add_slide(prs.slide_layouts[6])
set_background(slide)
add_title(slide, "Literature Review", "Slide 3")
add_card(
    slide,
    Inches(0.8),
    Inches(1.5),
    Inches(3.9),
    Inches(3.1),
    "Zero Trust",
    "Research emphasizes: never trust by default, verify explicitly, and apply least privilege.",
)
add_card(
    slide,
    Inches(4.95),
    Inches(1.5),
    Inches(3.9),
    Inches(3.1),
    "Adaptive Security",
    "Context-aware access uses user role, device state, and behavior to make decisions.",
)
add_card(
    slide,
    Inches(9.1),
    Inches(1.5),
    Inches(3.4),
    Inches(3.1),
    "MFA + Risk",
    "Combining MFA with risk scoring reduces unauthorized access chances.",
)

# Slide 4: Proposed approach + Zero Trust
slide = prs.slides.add_slide(prs.slide_layouts[6])
set_background(slide)
add_title(slide, "Proposed Approach and Zero Trust Aspect", "Slide 4")
add_bullets(
    slide,
    Inches(0.8),
    Inches(1.3),
    Inches(6.3),
    Inches(5.2),
    [
        "User login with role and OTP verification.",
        "Device posture and risk score evaluation.",
        "Policy decision: allow, challenge, or block.",
        "Continuous traffic monitoring and anomaly detection.",
        "Session logging for audit and replay.",
    ],
    size=19,
)
add_card(
    slide,
    Inches(7.4),
    Inches(1.6),
    Inches(5.2),
    Inches(1.5),
    "Never Trust",
    "Every session is validated using identity, device, and behavior signals.",
)
add_card(
    slide,
    Inches(7.4),
    Inches(3.35),
    Inches(5.2),
    Inches(1.5),
    "Least Privilege",
    "RBAC restricts access by role and prevents unsafe privilege usage.",
)
add_card(
    slide,
    Inches(7.4),
    Inches(5.1),
    Inches(5.2),
    Inches(1.5),
    "Assume Breach",
    "Suspicious sessions can be alerted, challenged, and blocked quickly.",
)

# Slide 5: Results + future work
slide = prs.slides.add_slide(prs.slide_layouts[6])
set_background(slide)
add_title(slide, "Results and Future Work", "Slide 5")
add_card(
    slide,
    Inches(0.8),
    Inches(1.35),
    Inches(5.9),
    Inches(2.2),
    "Current Results",
    "Working desktop GUI with login flow, risk-based access decision, active sessions, alerts, and traffic monitoring.",
)
add_card(
    slide,
    Inches(0.8),
    Inches(3.85),
    Inches(5.9),
    Inches(2.2),
    "Future Work",
    "Integrate stronger ML models, connect real identity provider, and move logs to a scalable database-backed backend.",
)
add_bullets(
    slide,
    Inches(7.1),
    Inches(1.5),
    Inches(5.4),
    Inches(4.2),
    [
        "Improve real-time threat analytics dashboard.",
        "Automate policy tuning from incident patterns.",
        "Deploy as production-grade enterprise architecture.",
    ],
    size=18,
)

thanks = slide.shapes.add_textbox(Inches(7.1), Inches(6.05), Inches(5.0), Inches(0.6))
ttf = thanks.text_frame
tp = ttf.paragraphs[0]
tr = tp.add_run()
tr.text = "Thank You"
tr.font.size = Pt(28)
tr.font.bold = True
tr.font.color.rgb = TEXT

prs.save(str(OUT_PATH))
print(f"Saved presentation to {OUT_PATH}")
